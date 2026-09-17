"""Extract plain text from an uploaded resume file.

Supports PDF, DOCX, and plain text. The extracted text is what gets sent to
Claude as the candidate's background, so it doesn't need to be pretty -- just
complete.
"""
from __future__ import annotations

import io
import os

from pypdf import PdfReader
import docx  # python-docx


class UnsupportedResumeFormat(ValueError):
    pass


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Return the plain-text contents of a resume file.

    Raises UnsupportedResumeFormat if the extension isn't one we handle.
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return _extract_pdf(file_bytes)
    if ext == ".docx":
        return _extract_docx(file_bytes)
    if ext in (".txt", ".md"):
        return file_bytes.decode("utf-8", errors="replace")

    raise UnsupportedResumeFormat(
        f"Unsupported resume format '{ext}'. Supported formats: "
        f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)
    text = "\n".join(pages).strip()
    if not text:
        raise UnsupportedResumeFormat(
            "Couldn't extract any text from this PDF -- it may be a scanned "
            "image rather than a text-based PDF. Try exporting a text-based "
            "PDF, or upload a .docx / .txt version instead."
        )
    return text


def _extract_docx(file_bytes: bytes) -> str:
    document = docx.Document(io.BytesIO(file_bytes))
    parts = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Tables (some resumes use them for skills/experience layout)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise UnsupportedResumeFormat("Couldn't find any text in this .docx file.")
    return text
