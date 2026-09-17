import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx
import pytest

from core.resume_parser import extract_text, UnsupportedResumeFormat


def test_extract_txt():
    text = extract_text(b"Hello resume", "resume.txt")
    assert text == "Hello resume"


def test_extract_md():
    text = extract_text(b"# Resume\n\nSome text", "resume.md")
    assert "Some text" in text


def test_extract_docx():
    buf = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph("Jane Smith")
    doc.add_paragraph("Senior Product Manager")
    doc.save(buf)
    text = extract_text(buf.getvalue(), "resume.docx")
    assert "Jane Smith" in text
    assert "Senior Product Manager" in text


def test_unsupported_extension():
    with pytest.raises(UnsupportedResumeFormat):
        extract_text(b"whatever", "resume.pages")
